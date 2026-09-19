#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
模拟数据生成脚本 - 用于测试「文件清理助手」
生成测试目录树和 Excel 清单
"""

import os
import random
import argparse
from pathlib import Path
import openpyxl
from openpyxl.styles import Font

# 预设名称池
PREFIXES = ['报告', '数据', '备份', '临时', '旧版', '草稿', '项目', '文档', '图片', '视频']
SUFFIXES = ['2023', '2024', 'final', 'v1', 'v2', 'old', 'temp', 'test', 'backup', 'draft']
EXTENSIONS = ['.txt', '.docx', '.xlsx', '.pdf', '.jpg', '.png', '.log', '.tmp']

SPECIAL_CHARS = [' ', '-', '_', '（', '）', '【', '】', '!', '@', '#']
LONG_NAME = '这是一个非常长的文件名用于测试超长路径和名称的处理能力' * 3

def random_name(include_special=False, include_long=False):
    """生成随机名称"""
    if include_long and random.random() < 0.05:
        return LONG_NAME[:200]  # 限制长度
    name = random.choice(PREFIXES) + random.choice(SUFFIXES)
    if include_special and random.random() < 0.3:
        name += random.choice(SPECIAL_CHARS)
    return name

def random_filename(include_special=False, include_long=False):
    """生成随机文件名（带扩展名）"""
    return random_name(include_special, include_long) + random.choice(EXTENSIONS)

def create_file(path, content='test content', readonly=False):
    """创建文件，可设置只读"""
    path.write_text(content, encoding='utf-8')
    if readonly:
        os.chmod(path, 0o444)  # 只读

def build_tree(root, depth, files_per_dir, dirs_per_dir,
               include_special, include_long, include_empty_dirs, include_readonly):
    """递归构建目录树，返回所有生成的文件名和文件夹名集合"""
    all_files = set()
    all_dirs = set()

    def _build(current_path, current_depth):
        if current_depth > depth:
            return
        # 创建文件
        for i in range(files_per_dir):
            fname = random_filename(include_special, include_long)
            fpath = current_path / fname
            counter = 1
            while fpath.exists():
                fpath = current_path / f"{fpath.stem}_{counter}{fpath.suffix}"
                counter += 1
            readonly = include_readonly and random.random() < 0.1
            create_file(fpath, f"内容 {fname}", readonly)
            all_files.add(fpath.name)
        # 创建子目录
        for i in range(dirs_per_dir):
            dname = random_name(include_special, include_long)
            dpath = current_path / dname
            counter = 1
            while dpath.exists():
                dpath = current_path / f"{dname}_{counter}"
                counter += 1
            dpath.mkdir(parents=True, exist_ok=True)
            all_dirs.add(dpath.name)
            # 空文件夹
            if include_empty_dirs and random.random() < 0.2:
                empty_dir = dpath / 'empty_folder'
                empty_dir.mkdir(exist_ok=True)
                all_dirs.add(empty_dir.name)
            _build(dpath, current_depth + 1)

    _build(root, 1)
    return all_files, all_dirs

def generate_excel(excel_path, all_files, all_dirs,
                   include_duplicates, include_blank, include_mismatch, match_keywords):
    """生成 Excel 清单"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "待删除清单"

    # 表头
    headers = ['名称', '编号', '备注']
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)

    # 收集要写入的名称
    names_to_write = []
    # 1. 精确匹配：部分实际存在的文件/文件夹
    sample_files = random.sample(list(all_files), min(10, len(all_files)))
    sample_dirs = random.sample(list(all_dirs), min(10, len(all_dirs)))
    names_to_write.extend(sample_files)
    names_to_write.extend(sample_dirs)

    # 2. 包含匹配关键词（这些关键词可能出现在文件名中）
    names_to_write.extend(match_keywords)

    # 3. 大小写变体
    if sample_files:
        variant = random.choice(sample_files).upper()
        names_to_write.append(variant)
        variant = random.choice(sample_files).lower()
        names_to_write.append(variant)

    # 4. 重复项
    if include_duplicates and names_to_write:
        names_to_write.append(random.choice(names_to_write))
        names_to_write.append(random.choice(names_to_write))

    # 5. 空行（空字符串）
    if include_blank:
        names_to_write.append('')
        names_to_write.append('   ')

    # 6. 不匹配的随机名称
    if include_mismatch:
        for _ in range(5):
            names_to_write.append(random_name() + '_不存在')

    # 7. 特殊字符名称（从实际文件/文件夹中筛选）
    special_names = [n for n in list(all_files) + list(all_dirs)
                     if any(c in n for c in SPECIAL_CHARS)]
    if special_names:
        names_to_write.extend(random.sample(special_names, min(3, len(special_names))))

    # 写入 Excel
    row = 2
    for name in names_to_write:
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=f"ID{row-1:04d}")
        ws.cell(row=row, column=3, value="测试数据")
        row += 1

    # 自适应列宽
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[column].width = min(max_length + 2, 50)

    wb.save(excel_path)
    return row - 2

def main():
    parser = argparse.ArgumentParser(description='生成文件清理助手的模拟测试数据')
    parser.add_argument('--root', type=str, default='./test_data', help='测试根目录')
    parser.add_argument('--excel', type=str, default='./test_data/test_list.xlsx', help='Excel输出路径')
    parser.add_argument('--depth', type=int, default=3, help='目录深度')
    parser.add_argument('--files-per-dir', type=int, default=5, help='每个目录下文件数')
    parser.add_argument('--dirs-per-dir', type=int, default=2, help='每个目录下子目录数')
    parser.add_argument('--include-special', action='store_true', default=True, help='包含特殊字符')
    parser.add_argument('--no-special', action='store_false', dest='include_special', help='不包含特殊字符')
    parser.add_argument('--include-long', action='store_true', default=True, help='包含长文件名')
    parser.add_argument('--no-long', action='store_false', dest='include_long', help='不包含长文件名')
    parser.add_argument('--include-empty-dirs', action='store_true', default=True, help='包含空文件夹')
    parser.add_argument('--no-empty-dirs', action='store_false', dest='include_empty_dirs', help='不包含空文件夹')
    parser.add_argument('--include-readonly', action='store_true', default=True, help='包含只读文件')
    parser.add_argument('--no-readonly', action='store_false', dest='include_readonly', help='不包含只读文件')
    parser.add_argument('--include-duplicates', action='store_true', default=True, help='Excel中包含重复项')
    parser.add_argument('--no-duplicates', action='store_false', dest='include_duplicates', help='不包含重复项')
    parser.add_argument('--include-blank', action='store_true', default=True, help='Excel中包含空行')
    parser.add_argument('--no-blank', action='store_false', dest='include_blank', help='不包含空行')
    parser.add_argument('--include-mismatch', action='store_true', default=True, help='包含不匹配的项')
    parser.add_argument('--no-mismatch', action='store_false', dest='include_mismatch', help='不包含不匹配项')
    parser.add_argument('--keywords', type=str, default='temp,old,test,backup',
                        help='包含匹配的关键词，逗号分隔')

    args = parser.parse_args()

    root = Path(args.root).resolve()
    excel_path = Path(args.excel).resolve()

    # 清理旧数据
    if root.exists():
        import shutil
        print(f"警告：根目录 {root} 已存在，将清空并重建。")
        shutil.rmtree(root)

    root.mkdir(parents=True, exist_ok=True)
    excel_path.parent.mkdir(parents=True, exist_ok=True)

    print(f"开始生成测试数据...")
    print(f"  根目录: {root}")
    print(f"  Excel: {excel_path}")

    all_files, all_dirs = build_tree(
        root,
        depth=args.depth,
        files_per_dir=args.files_per_dir,
        dirs_per_dir=args.dirs_per_dir,
        include_special=args.include_special,
        include_long=args.include_long,
        include_empty_dirs=args.include_empty_dirs,
        include_readonly=args.include_readonly
    )

    print(f"  生成文件: {len(all_files)} 个")
    print(f"  生成文件夹: {len(all_dirs)} 个")

    keywords = [k.strip() for k in args.keywords.split(',') if k.strip()]
    rows = generate_excel(
        excel_path,
        all_files,
        all_dirs,
        include_duplicates=args.include_duplicates,
        include_blank=args.include_blank,
        include_mismatch=args.include_mismatch,
        match_keywords=keywords
    )

    print(f"  Excel 写入: {rows} 行数据")
    print("生成完成！")
    print("\n使用说明：")
    print(f"1. 打开「文件清理助手」，选择 Excel 文件: {excel_path}")
    print(f"2. 选择目标路径: {root}")
    print("3. 根据需要配置删除类型、匹配模式，点击「扫描预览」")
    print("4. 测试完成后，可重新运行本脚本重置数据")

if __name__ == '__main__':
    main()